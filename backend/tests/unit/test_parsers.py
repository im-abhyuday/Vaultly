"""Unit tests for document parsers (PDF, DOCX, TXT/MD)."""

import io
from pathlib import Path

import docx
import pypdf
import pytest

from app.models.enums import DocumentErrorCode
from app.services.parser import (
    clean_text,
    parse_docx_stream,
    parse_pdf_stream,
    parse_text_stream,
    sync_parse_document,
)
from app.services.storage import LocalStorage
from app.workers.errors import PermanentError


def create_in_memory_pdf(pages_text: list[str]) -> io.BytesIO:
    """Create a synthetic in-memory PDF using pypdf."""
    writer = pypdf.PdfWriter()
    for _ in pages_text:
        # Create a blank page
        writer.add_blank_page(width=612, height=792)

        # Note: blank pages in pypdf don't have text stream by default unless added
        # We can add an annotation or text or mock extract_text if needed
    stream = io.BytesIO()
    writer.write(stream)
    stream.seek(0)
    return stream


def test_clean_text_nul_bytes_and_whitespace() -> None:
    """clean_text strips NUL bytes, normalizes newlines, and collapses whitespace."""
    raw = "Header\x00 text\r\n\r\n\r\n\r\nBody   with    spaces\x00 and\ttabs."
    cleaned = clean_text(raw)
    assert "\x00" not in cleaned
    assert "\r" not in cleaned
    assert "    " not in cleaned
    assert cleaned == "Header text\n\nBody with spaces and tabs."


def test_parse_text_stream_utf8() -> None:
    """parse_text_stream decodes UTF-8 and UTF-8-SIG cleanly."""
    content = "Vaultly enterprise RAG platform.\nSecure multi-tenant data."
    raw_utf8 = content.encode("utf-8")
    segments, page_count = parse_text_stream(raw_utf8)
    assert page_count == 1
    assert len(segments) == 1
    assert segments[0].page == 1
    assert "Vaultly enterprise RAG" in segments[0].text

    # UTF-8 with BOM
    raw_bom = content.encode("utf-8-sig")
    segments_bom, _ = parse_text_stream(raw_bom)
    assert segments_bom[0].text == segments[0].text


def test_parse_text_stream_cp1252() -> None:
    """parse_text_stream falls back to CP1252 for legacy Windows encodings."""
    # Byte 0x93 and 0x94 are smart quotes in cp1252, invalid in utf-8
    cp1252_bytes = b"Vaultly \x93Secure\x94 storage"
    segments, page_count = parse_text_stream(cp1252_bytes)
    assert page_count == 1
    assert "Vaultly" in segments[0].text
    assert "Secure" in segments[0].text


def test_parse_text_stream_empty_raises_no_text() -> None:
    """parse_text_stream on whitespace-only input raises NO_EXTRACTABLE_TEXT."""
    with pytest.raises(PermanentError) as exc_info:
        parse_text_stream(b"   \r\n   \t  ")
    assert exc_info.value.error_code == DocumentErrorCode.NO_EXTRACTABLE_TEXT.value


def test_parse_pdf_stream_success(monkeypatch: pytest.MonkeyPatch) -> None:
    """parse_pdf_stream extracts per-page text segments."""
    stream = create_in_memory_pdf(["Page 1", "Page 2"])

    # Mock extract_text on pages since pypdf blank pages don't render text
    page_texts = ["Page 1: Vaultly overview.", "Page 2: Security and RLS architecture."]

    def mock_extract(_self: pypdf.PageObject) -> str:
        # Return page_texts based on some counter
        return page_texts.pop(0) if page_texts else "Extra page"


    monkeypatch.setattr(pypdf.PageObject, "extract_text", mock_extract)

    segments, count = parse_pdf_stream(stream, max_pages=10)
    assert count == 2
    assert len(segments) == 2
    assert segments[0].page == 1
    assert "Vaultly overview" in segments[0].text
    assert segments[1].page == 2
    assert "Security and RLS" in segments[1].text


def test_parse_pdf_stream_corrupt_file() -> None:
    """Corrupt PDF bytes raise PermanentError CORRUPT_FILE."""
    corrupt_stream = io.BytesIO(b"%PDF-1.4 corrupt content not a real pdf %%EOF")
    with pytest.raises(PermanentError) as exc_info:
        parse_pdf_stream(corrupt_stream, max_pages=10)
    assert exc_info.value.error_code == DocumentErrorCode.CORRUPT_FILE.value


def test_parse_pdf_stream_page_limit_exceeded() -> None:
    """PDF with more pages than max_pages raises PermanentError LIMIT_EXCEEDED."""
    stream = create_in_memory_pdf(["P1", "P2", "P3", "P4", "P5"])
    with pytest.raises(PermanentError) as exc_info:
        parse_pdf_stream(stream, max_pages=3)
    assert exc_info.value.error_code == DocumentErrorCode.LIMIT_EXCEEDED.value


def test_parse_pdf_stream_scanned_image_only(monkeypatch: pytest.MonkeyPatch) -> None:
    """PDF with zero extractable text raises PermanentError NO_EXTRACTABLE_TEXT."""
    stream = create_in_memory_pdf(["Blank"])
    monkeypatch.setattr(pypdf.PageObject, "extract_text", lambda _self: "")


    with pytest.raises(PermanentError) as exc_info:
        parse_pdf_stream(stream, max_pages=10)
    assert exc_info.value.error_code == DocumentErrorCode.NO_EXTRACTABLE_TEXT.value


def test_parse_docx_stream_paragraphs_and_tables() -> None:
    """parse_docx_stream extracts paragraphs and tabular data."""
    doc = docx.Document()
    doc.add_heading("Vaultly Architecture", level=1)
    doc.add_paragraph("This is a paragraph detailing multi-tenant RAG.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).text = "PostgreSQL"
    table.cell(1, 1).text = "Database with RLS"

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    segments, page_count = parse_docx_stream(stream)
    assert page_count == 1
    assert len(segments) == 1
    text = segments[0].text
    assert "Vaultly Architecture" in text
    assert "This is a paragraph" in text
    assert "Component | Role" in text
    assert "PostgreSQL | Database with RLS" in text


def test_parse_docx_stream_corrupt_file() -> None:
    """Invalid DOCX bytes raise PermanentError CORRUPT_FILE."""
    corrupt_stream = io.BytesIO(b"PK\x03\x04 corrupt docx zip stream")
    with pytest.raises(PermanentError) as exc_info:
        parse_docx_stream(corrupt_stream)
    assert exc_info.value.error_code == DocumentErrorCode.CORRUPT_FILE.value


def test_parse_docx_stream_empty_raises_no_text() -> None:
    """Empty DOCX without paragraphs or tables raises PermanentError NO_EXTRACTABLE_TEXT."""
    doc = docx.Document()
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    with pytest.raises(PermanentError) as exc_info:
        parse_docx_stream(stream)
    assert exc_info.value.error_code == DocumentErrorCode.NO_EXTRACTABLE_TEXT.value


def test_sync_parse_document_extracted_chars_limit(tmp_path: Path) -> None:
    """Extracted text exceeding max_extracted_chars raises PermanentError LIMIT_EXCEEDED."""
    storage = LocalStorage(str(tmp_path))
    content = "A" * 10_000
    storage_key = "test_doc.txt"
    (storage.root_path / storage_key).write_bytes(content.encode("utf-8"))


    with pytest.raises(PermanentError) as exc_info:
        sync_parse_document(
            storage=storage,
            storage_key=storage_key,
            detected_type="txt",
            max_extracted_chars=5000,
        )
    assert exc_info.value.error_code == DocumentErrorCode.LIMIT_EXCEEDED.value


def test_sync_parse_document_missing_file(tmp_path: Path) -> None:
    """Missing storage key raises PermanentError CORRUPT_FILE."""
    storage = LocalStorage(str(tmp_path))
    with pytest.raises(PermanentError) as exc_info:
        sync_parse_document(
            storage=storage,
            storage_key="nonexistent.txt",
            detected_type="txt",
        )
    assert exc_info.value.error_code == DocumentErrorCode.CORRUPT_FILE.value
